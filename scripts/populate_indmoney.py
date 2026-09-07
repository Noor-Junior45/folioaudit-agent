import os
import sys
import openpyxl
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

GROWW_DIR = r"c:\Users\mdhas\Downloads\Groww"
TEMPLATE_PATH = os.path.join(GROWW_DIR, "indmoney_template.xlsx")
STOCKS_REPORT_PATH = os.path.join(GROWW_DIR, "Stocks_Capital_Gains_Report_1723632298_01-04-2026_31-03-2027.xlsx")
DIVIDENDS_PATH = os.path.join(GROWW_DIR, "dividends-QEJ392-2026_2027.xlsx")

def format_date_str(val):
    if not val:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%d/%m/%Y")
    val_str = str(val).strip()
    # If YYYY-MM-DD
    if len(val_str) == 10 and val_str[4] == "-" and val_str[7] == "-":
        parts = val_str.split("-")
        return f"{parts[2]}/{parts[1]}/{parts[0]}"
    # If DD-MM-YYYY
    if len(val_str) == 10 and val_str[2] == "-" and val_str[5] == "-":
        parts = val_str.split("-")
        return f"{parts[0]}/{parts[1]}/{parts[2]}"
    # If already DD/MM/YYYY
    if "/" in val_str:
        return val_str
    return val_str

def populate_indmoney():
    print("Loading indmoney template...")
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    
    # 1. Populate Indian Stocks
    print("Reading Stocks Capital Gains Report...")
    wb_stocks = openpyxl.load_workbook(STOCKS_REPORT_PATH, data_only=True)
    ws_src = wb_stocks["Sheet1"]
    
    stock_trades = []
    # Short term (rows 38 to 42)
    for r in range(38, 43):
        name = ws_src.cell(r, 1).value
        if name:
            stock_trades.append({
                "name": str(name).strip(),
                "isin": str(ws_src.cell(r, 2).value).strip(),
                "qty": float(ws_src.cell(r, 3).value or 0),
                "buy_date": format_date_str(ws_src.cell(r, 4).value),
                "buy_val": float(ws_src.cell(r, 6).value or 0),
                "sell_date": format_date_str(ws_src.cell(r, 7).value),
                "sell_val": float(ws_src.cell(r, 9).value or 0),
            })
            
    # Long term (rows 46 to 78)
    for r in range(46, 79):
        name = ws_src.cell(r, 1).value
        if name:
            stock_trades.append({
                "name": str(name).strip(),
                "isin": str(ws_src.cell(r, 2).value).strip(),
                "qty": float(ws_src.cell(r, 3).value or 0),
                "buy_date": format_date_str(ws_src.cell(r, 4).value),
                "buy_val": float(ws_src.cell(r, 6).value or 0),
                "sell_date": format_date_str(ws_src.cell(r, 7).value),
                "sell_val": float(ws_src.cell(r, 9).value or 0),
            })
    wb_stocks.close()
    print(f"Extracted {len(stock_trades)} stock trades.")
    
    ws_stocks = wb["📈 Indian Stocks"]
    for idx, trade in enumerate(stock_trades):
        row_num = 4 + idx
        ws_stocks.cell(row=row_num, column=1).value = trade["name"]
        ws_stocks.cell(row=row_num, column=2).value = trade["isin"]
        ws_stocks.cell(row=row_num, column=3).value = trade["buy_date"]
        ws_stocks.cell(row=row_num, column=4).value = trade["sell_date"]
        ws_stocks.cell(row=row_num, column=5).value = trade["qty"]
        ws_stocks.cell(row=row_num, column=6).value = trade["buy_val"]
        ws_stocks.cell(row=row_num, column=7).value = trade["sell_val"]
        ws_stocks.cell(row=row_num, column=8).value = 0.0 # Expenses
        ws_stocks.cell(row=row_num, column=9).value = 0.0 # STT
    print(f"Populated {len(stock_trades)} trades into '📈 Indian Stocks'.")

    # 2. Populate Dividends
    print("Reading Dividends Report...")
    wb_div = openpyxl.load_workbook(DIVIDENDS_PATH, data_only=True)
    ws_div_src = wb_div["Equity Dividends"]
    
    dividends = []
    for r in range(16, 22):
        sym = ws_div_src.cell(r, 2).value
        if sym:
            sym_clean = str(sym).replace("#", "").strip()
            dividends.append({
                "name": sym_clean,
                "isin": str(ws_div_src.cell(r, 3).value).strip(),
                "type": "Stock",
                "date": format_date_str(ws_div_src.cell(r, 4).value),
                "amount": float(ws_div_src.cell(r, 7).value or 0),
            })
    wb_div.close()
    print(f"Extracted {len(dividends)} dividend records.")
    
    ws_div = wb["💰 Dividends"]
    for idx, div in enumerate(dividends):
        row_num = 4 + idx
        ws_div.cell(row=row_num, column=1).value = div["name"]
        ws_div.cell(row=row_num, column=2).value = div["isin"]
        ws_div.cell(row=row_num, column=3).value = div["type"]
        ws_div.cell(row=row_num, column=4).value = div["date"]
        ws_div.cell(row=row_num, column=5).value = div["amount"]
    print(f"Populated {len(dividends)} records into '💰 Dividends'.")

    # Save template
    print(f"Saving updated template to {TEMPLATE_PATH}...")
    wb.save(TEMPLATE_PATH)
    wb.close()
    print("SUCCESS: indmoney_template.xlsx successfully populated with Groww data!")

if __name__ == "__main__":
    populate_indmoney()
